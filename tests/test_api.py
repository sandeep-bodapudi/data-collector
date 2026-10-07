"""Backend tests. Run:  python tests/test_api.py   (exit code 1 if anything fails). Uses a throwaway SQLite database."""
import base64
import io
import json
import os
import sys
import tempfile
import time

TMP = tempfile.mkdtemp()
os.environ["DATABASE_URL"] = "sqlite:///" + os.path.join(TMP, "t.db").replace("\\", "/")
os.environ["APP_USER"] = "boss"
os.environ["APP_PASSWORD"] = "adminpass123"
os.environ["SECRET_KEY"] = "test-secret-key-for-unit-tests-only"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

import app as appmod  # noqa: E402
from models import VaultCredential  # noqa: E402
from scraper import jobs as J  # noqa: E402

app = appmod.app
failures = []


def ok(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail and not cond else ""))
    if not cond:
        failures.append(name)


c = app.test_client()

# ---- public PWA files ------------------------------------------------------------------------
m = c.get("/manifest.webmanifest")
mj = m.get_json(force=True)
ok("manifest is public and valid", m.status_code == 200 and mj["display"] == "standalone" and mj["start_url"] == "/")
sizes = {(i["sizes"], i.get("purpose")) for i in mj["icons"]}
ok("manifest has 192, 512 and maskable icons", {("192x192", "any"), ("512x512", "any"), ("512x512", "maskable")} <= sizes)
ok("every manifest icon file exists", all(c.get(i["src"]).status_code == 200 for i in mj["icons"]))
sw = c.get("/sw.js")
ok("service worker served from root, never cached", sw.status_code == 200 and "no-cache" in sw.headers["Cache-Control"]
   and sw.headers["Service-Worker-Allowed"] == "/" and "javascript" in sw.mimetype)
ok("service worker precaches the app files", b"/static/app.js?v=" in sw.data and b"/offline.html" in sw.data)
ok("service worker never caches the API", b'startsWith("/api/")' in sw.data)
ok("offline page is public", c.get("/offline.html").status_code == 200)
login_page = c.get("/login")
ok("login page links manifest and logo", b"/manifest.webmanifest" in login_page.data and b"logo.svg" in login_page.data)
ok("logo is a real SVG", b"<svg" in c.get("/static/logo.svg").data[:200])

# ---- access control --------------------------------------------------------------------------
ok("home redirects to login", c.get("/").status_code == 302)
ok("api gives 401 when signed out", c.get("/api/settings").status_code == 401)
ok("old server-side sheet routes are gone", c.get("/api/sheets").status_code in (401, 404))
r = c.post("/login", data={"username": "BOSS", "password": "adminpass123"})
ok("admin from APP_USER/APP_PASSWORD can sign in", r.status_code == 302)
home = c.get("/")
ok("app shell renders with PWA tags and logo", home.status_code == 200 and b'rel="manifest"' in home.data and b"brand-logo" in home.data
   and b"store.js?v=" in home.data and b"sheetops.js?v=" in home.data)
ok("app config includes the user id (used to name the browser database)", b'"id": 1' in home.data or b'"id":1' in home.data)

# ---- settings: keys encrypted, never returned ------------------------------------------------
c.post("/api/settings", json={"ai_provider": "openai", "ai_api_key": "sk-test-1234567890ABCD"})
s = c.get("/api/settings").get_json()
ok("saved key is masked in the API", s["ai_key_mask"].endswith("ABCD") and "sk-test" not in json.dumps(s))
with app.app_context():
    raw = VaultCredential.query.filter_by(name="ai_api_key").first().encrypted_value
ok("saved key is encrypted in the database", "sk-test" not in raw)
c.post("/api/settings", json={"brave_api_key": "BSA-test-key-0123456789"})
bs = c.get("/api/settings").get_json()
ok("a Brave key is saved to this user's account and shown only masked",
   bs["brave_key_mask"].endswith("6789") and "BSA-test" not in json.dumps(bs))
ok("testing the key with no key saved says so, without calling Brave",
   app.test_client().post("/api/settings/brave/test").status_code == 401)
c.post("/api/settings", json={"brave_api_key": ""})
ok("clearing the key removes it", c.get("/api/settings").get_json()["brave_key_mask"] == "")
ok("cookies are ignored while restricted mode is off", c.post("/api/settings", json={"li_at_cookie": "x"}).status_code == 200
   and c.get("/api/settings").get_json()["li_cookie_mask"] == "")

# ---- runs: live on the server only until the browser collects them ---------------------------
def fake_web(job):
    job.columns = ["Name", "Emails"]
    job.rows = [{"Name": "A", "Emails": ["a@a.com", "b@b.com"]}, {"Name": "B", "Emails": []}]
    job.total = job.done = 2
J._run_web = fake_web
r = c.post("/api/jobs", json={"mode": "web", "queries": "schools", "fields": ["emails"]})
jid = r.get_json()["id"]
for _ in range(50):
    if c.get(f"/api/jobs/{jid}").get_json()["status"] not in ("queued", "running"):
        break
    time.sleep(0.1)
st = c.get(f"/api/jobs/{jid}").get_json()
ok("run finishes", st["status"] == "done" and st["count"] == 2, st["status"])
res = c.get(f"/api/jobs/{jid}/result").get_json()
ok("result gives plain-text cells (lists joined)", res["rows"][0]["Emails"] == "a@a.com; b@b.com" and res["columns"] == ["Name", "Emails"])
c.post("/api/admin/users", json={"username": "ravi", "password": "memberpass1", "role": "viewer"})
c.post("/api/admin/users", json={"username": "meena", "password": "memberpass2", "role": "member"})
meena = app.test_client(); meena.post("/login", data={"username": "meena", "password": "memberpass2"})
ok("another user cannot see or fetch my run", meena.get(f"/api/jobs/{jid}").status_code == 404 and meena.get(f"/api/jobs/{jid}/result").status_code == 404)
ok("another user cannot delete my run", meena.delete(f"/api/jobs/{jid}").status_code == 404)
ok("forgetting a run frees the server's copy", c.delete(f"/api/jobs/{jid}").status_code == 200 and c.get(f"/api/jobs/{jid}").status_code == 404)

# a run that is stopped while running
def slow_web(job):
    job.columns = ["Name"]
    for _ in range(100):
        if job.cancelled.is_set():
            return
        time.sleep(0.05)
J._run_web = slow_web
jid2 = c.post("/api/jobs", json={"mode": "web", "queries": "x", "fields": ["emails"]}).get_json()["id"]
ok("result of a running job is refused (409)", c.get(f"/api/jobs/{jid2}/result").status_code == 409)
ok("a running job cannot be forgotten (409)", c.delete(f"/api/jobs/{jid2}").status_code == 409)
c.post(f"/api/jobs/{jid2}/stop")
for _ in range(50):
    if c.get(f"/api/jobs/{jid2}").get_json()["status"] == "cancelled":
        break
    time.sleep(0.1)
ok("stopping a run works", c.get(f"/api/jobs/{jid2}").get_json()["status"] == "cancelled")

# outage => failed with the reason
def boom(*a, **k):
    raise J.places.PlaceError("OpenStreetMap servers are busy (HTTP 504). Please try again in a few minutes.")
J.places.search_places = boom
J.time.sleep = lambda s: None
job = J.Job({"mode": "places", "category": "Hindu temples", "locations": ["X"], "max_results": 5, "custom_fields": [], "file_name": "", "ai": {}}, 1)
J._run(job)
ok("a service outage shows as failed with the reason", job.status == "error" and "busy" in job.error)
old = J.Job({"mode": "web"}, 1); old.finished = J.datetime(2000, 1, 1); J.JOBS["old"] = old; J.purge_jobs()
ok("purge removes expired finished jobs", "old" not in J.JOBS)

# viewer cannot start runs
viewer = app.test_client(); viewer.post("/login", data={"username": "ravi", "password": "memberpass1"})
ok("viewer cannot start runs", viewer.post("/api/jobs", json={"mode": "web", "queries": "x", "fields": ["emails"]}).status_code == 403)
ok("viewer cannot use admin api", viewer.get("/api/admin/users").status_code == 403)
ok("viewer can still export a sheet they hold", viewer.post("/api/export/xlsx", json={"name": "n", "columns": ["a"], "rows": [{"a": 1}]}).status_code == 200)

# ---- stateless helpers -----------------------------------------------------------------------
csv_file = (io.BytesIO(b"Name,Phone\nA,123\nB,456\n"), "one.csv")
r = c.post("/api/parse", data={"file": csv_file}, content_type="multipart/form-data").get_json()
ok("parse reads a CSV", r["columns"] == ["Name", "Phone"] and len(r["rows"]) == 2 and r["name"] == "one")
ok("parse rejects other file types", c.post("/api/parse", data={"file": (io.BytesIO(b"x"), "a.exe")}, content_type="multipart/form-data").status_code == 400)
ok("parse rejects an empty request", c.post("/api/parse", data={}, content_type="multipart/form-data").status_code == 400)
ok("viewer cannot import", viewer.post("/api/parse", data={"file": (io.BytesIO(b"a\n1\n"), "a.csv")}, content_type="multipart/form-data").status_code == 403)

x = c.post("/api/export/xlsx", json={"name": "My sheet", "columns": ["Name", "Site"],
                                     "rows": [{"Name": "=HYPERLINK(\"http://evil\",\"x\")", "Site": "https://a.com"}]})
wb = load_workbook(io.BytesIO(x.data))
ok("xlsx export is a real workbook", x.data[:2] == b"PK" and wb.sheetnames == ["Data", "Run Info"])
ok("scraped text starting with = stays text", wb["Data"]["A2"].data_type == "s")
ok("xlsx export rejects an empty request", c.post("/api/export/xlsx", json={}).status_code == 400)

# ---- admin ----------------------------------------------------------------------------------
a = c.get("/api/admin/settings").get_json()
ok("admin sees database status", a["database"]["kind"] == "sqlite" and a["database"]["persistent"] is False)
ok("cannot demote the last admin", c.patch("/api/admin/users/1", json={"role": "member"}).status_code == 400)

# ---- sharing -----------------------------------------------------------------------------------
from datetime import datetime as _dt, timedelta as _td  # noqa: E402
from models import Share, db as _db  # noqa: E402

with app.app_context():
    from models import User as _U
    ids = {u.username: u.id for u in _U.query.all()}
ROWS = [{"Name": "శ్రీ School", "Phone": "+91 98480 12345", "n": 5}, {"Name": "B", "Phone": "", "n": 7}]
payload = lambda **kw: {"name": "My list", "columns": ["Name", "Phone", "n"], "rows": ROWS, **kw}

ok("viewer cannot share", viewer.post("/api/shares", json=payload(link_access=True)).status_code == 403)
ok("viewer cannot list people", viewer.get("/api/users/directory").status_code == 403)
names = [u["username"] for u in c.get("/api/users/directory").get_json()]
ok("people list excludes yourself", "boss" not in names and {"ravi", "meena"} <= set(names))
ok("sharing needs someone to share with", c.post("/api/shares", json=payload()).status_code == 400)
ok("expiry must be one of the offered choices", c.post("/api/shares", json=payload(link_access=True, expires_days=400)).status_code == 400)
ok("empty sheets are refused", c.post("/api/shares", json={"name": "x", "columns": ["a"], "rows": [], "link_access": True}).status_code == 400)
ok("unknown people are ignored, so no recipients is refused", c.post("/api/shares", json=payload(recipients=[999, 12345])).status_code == 400)

r = c.post("/api/shares", json=payload(recipients=[ids["ravi"]], allow_export=False, expires_days=7))
sh = r.get_json(); sid = sh["id"]
ok("owner can share with a person", r.status_code == 200 and sh["recipients"] == ["ravi"] and sh["allow_export"] is False and len(sid) >= 12)
opened = viewer.get(f"/api/shares/{sid}").get_json()
ok("recipient can open it", opened["rows_data"] == ROWS and opened["columns"] == ["Name", "Phone", "n"], str(opened)[:120])
ok("Telugu text and numbers survive the round trip", opened["rows_data"][0]["Name"] == "శ్రీ School" and opened["rows_data"][1]["n"] == 7)
ok("download setting is passed on to the recipient", opened["allow_export"] is False and opened["owner"] == "boss")
ok("someone who was not chosen cannot open it", meena.get(f"/api/shares/{sid}").status_code == 404)
ok("signed-out users cannot open it", app.test_client().get(f"/api/shares/{sid}").status_code == 401)
ok("owner always may export their own share", c.get(f"/api/shares/{sid}").get_json()["allow_export"] is True)
ok("it appears in the recipient's Shared with me", [x["id"] for x in viewer.get("/api/shares").get_json()] == [sid])
ok("it appears in the owner's Shared by me, not in with-me", [x["id"] for x in c.get("/api/shares?box=by-me").get_json()] == [sid]
   and c.get("/api/shares").get_json() == [])
ok("another member does not see it listed", meena.get("/api/shares").get_json() == [])

link = c.post("/api/shares", json=payload(link_access=True, expires_days=1)).get_json()
ok("anyone signed in with the link can open a link share", meena.get(f"/api/shares/{link['id']}").status_code == 200)
ok("link shares are not listed for people who were not chosen", all(x["id"] != link["id"] for x in meena.get("/api/shares").get_json()))
ok("link access still needs a sign-in", app.test_client().get(f"/api/shares/{link['id']}").status_code == 401)

ok("only the owner (or an admin) can stop a share", meena.delete(f"/api/shares/{link['id']}").status_code == 404)
ok("owner stops sharing", c.delete(f"/api/shares/{link['id']}").status_code == 200 and meena.get(f"/api/shares/{link['id']}").status_code == 404)
other = meena.post("/api/shares", json=payload(link_access=True)).get_json()
ok("an admin can stop anyone's share", c.delete(f"/api/shares/{other['id']}").status_code == 200)

with app.app_context():
    row = _db.session.get(Share, sid); row.expires_at = _dt.utcnow() - _td(minutes=1); _db.session.commit()
ok("expired shares cannot be opened", viewer.get(f"/api/shares/{sid}").status_code == 404)
ok("expired shares disappear from the list", viewer.get("/api/shares").get_json() == [])
with app.app_context():
    ok("expired shares are deleted from the database", _db.session.get(Share, sid) is None)

old_max = appmod.SHARE_MAX_BYTES; appmod.SHARE_MAX_BYTES = 50
big = c.post("/api/shares", json=payload(link_access=True, rows=[{"Name": "x" * 500 + str(i), "Phone": str(i), "n": i} for i in range(200)]))
ok("a sheet over the size limit is refused with advice", big.status_code == 413 and "too large" in big.get_json()["error"])
appmod.SHARE_MAX_BYTES = old_max
old_total = appmod.SHARE_MAX_TOTAL; appmod.SHARE_MAX_TOTAL = 10
ok("a full share space is reported clearly", c.post("/api/shares", json=payload(link_access=True)).status_code == 507)
appmod.SHARE_MAX_TOTAL = old_total
old_active = appmod.SHARE_MAX_ACTIVE; appmod.SHARE_MAX_ACTIVE = 1
first = c.post("/api/shares", json=payload(link_access=True))
ok("the number of active shares per person is capped", first.status_code == 200 and c.post("/api/shares", json=payload(link_access=True)).status_code == 400)
appmod.SHARE_MAX_ACTIVE = old_active

# ---- customer links (public pages for people outside the company) --------------------------------
CROWS = [{"Name": "శ్రీ Academy", "Email": "info@academy.in", "Note": "=HYPERLINK(\"http://evil\",\"x\")", "Phone": "+91 98480 12345"},
         {"Name": "B College", "Email": "", "Note": "ok", "Phone": "040-2345678"}]
cpay = lambda **kw: {"kind": "external", "name": "Student leads", "columns": ["Name", "Email", "Note", "Phone"], "rows": CROWS,
                     "customer": "Global Education Services", "acknowledged": True, "expires_days": 7, **kw}
ok("customer link needs the customer's name", c.post("/api/shares", json=cpay(customer="  ")).status_code == 400)
ok("customer link needs the confirmation", c.post("/api/shares", json=cpay(acknowledged=False)).status_code == 400)
ok("passcode must be 4 to 40 characters", c.post("/api/shares", json=cpay(passcode="abc")).status_code == 400)
ok("viewer cannot create a customer link", viewer.post("/api/shares", json=cpay()).status_code == 403)

created = c.post("/api/shares", json=cpay(passcode="Hyd2026", message="Hello,\nhere is the data you asked for."))
ext = created.get_json(); eid = ext["id"]
ok("customer link is created", created.status_code == 200 and ext["kind"] == "external" and ext["customer"] == "Global Education Services")
ok("response never contains the data or the passcode", "rows_data" not in ext and "Hyd2026" not in json.dumps(ext)
   and "passcode_hash" not in ext and ext["has_passcode"] is True)

ok("the sender can look their passcode up again", c.get(f"/api/shares/{eid}/passcode").get_json() == {"has_passcode": True, "passcode": "Hyd2026"})
ok("no one else can see the passcode", meena.get(f"/api/shares/{eid}/passcode").status_code == 404 and app.test_client().get(f"/api/shares/{eid}/passcode").status_code == 401)
with app.app_context():
    stored = _db.session.get(Share, eid)
    ok("the passcode is encrypted in the database, not plain text", stored.passcode_enc and "Hyd2026" not in stored.passcode_enc and "Hyd2026" not in stored.passcode_hash)
ok("share lists never carry the passcode", "Hyd2026" not in json.dumps(c.get("/api/shares?box=by-me").get_json()) and "Hyd2026" not in json.dumps(c.get("/api/admin/shares").get_json()))
ok("a link without a passcode has none to show", c.get(f"/api/shares/{c.post('/api/shares', json=cpay(name='Plain', customer='X')).get_json()['id']}/passcode").get_json() == {"has_passcode": False, "passcode": None})

anon = app.test_client()
H = {"X-Forwarded-For": "203.0.113.5"}
locked = anon.get(f"/s/{eid}", headers=H)
ok("the public page opens without signing in", locked.status_code == 200)
ok("before the passcode, no data and no sheet name is shown", b"protected" in locked.data and "Academy".encode() not in locked.data
   and b"Student leads" not in locked.data and "Global Education Services".encode() in locked.data)
ok("a protected page counts nothing before it is unlocked", c.get("/api/shares?box=by-me").get_json()[0]["views"] == 0)
ok("downloads need the passcode first", anon.get(f"/s/{eid}/file.csv", headers=H).status_code == 302)
wrong = anon.post(f"/s/{eid}", data={"passcode": "nope"}, headers=H)
ok("a wrong passcode is refused kindly", wrong.status_code == 200 and b"isn&#39;t right" in wrong.data and b"Academy" not in wrong.data)
for _ in range(5):
    last = anon.post(f"/s/{eid}", data={"passcode": "guess"}, headers=H)
ok("repeated wrong passcodes are locked out", b"Too many wrong attempts" in last.data)
ok("even the right passcode waits during the lockout", b"Too many wrong attempts" in anon.post(f"/s/{eid}", data={"passcode": "Hyd2026"}, headers=H).data)

cust = app.test_client()
H2 = {"X-Forwarded-For": "198.51.100.7"}
r = cust.post(f"/s/{eid}", data={"passcode": "Hyd2026"}, headers=H2)
ok("the right passcode opens the file", r.status_code == 302)
page = cust.get(f"/s/{eid}", headers=H2)
ok("the customer sees their data, the sender's note and the company", page.status_code == 200 and "శ్రీ Academy".encode() in page.data
   and b"Student leads" in page.data and b"Global Education Services" in page.data and b"here is the data" in page.data
   and b"OneBridge Infotech" in page.data)
ok("page is private: no caching, no indexing, no referrer", page.headers["Cache-Control"] == "no-store"
   and "noindex" in page.headers["X-Robots-Tag"] and page.headers["Referrer-Policy"] == "no-referrer")

csvr = cust.get(f"/s/{eid}/file.csv", headers=H2)
text = csvr.data.decode("utf-8-sig")
ok("CSV download works with a byte-order mark", csvr.status_code == 200 and csvr.data[:3] == b"\xef\xbb\xbf"
   and "attachment" in csvr.headers["Content-Disposition"])
ok("CSV keeps Telugu text and phone numbers", "శ్రీ Academy" in text and "+91 98480 12345" in text)
ok("CSV neutralises a formula but not a phone number", "'=HYPERLINK" in text and "'+91" not in text)
xr = cust.get(f"/s/{eid}/file.xlsx", headers=H2)
wbx = load_workbook(io.BytesIO(xr.data))
info = {row[0].value: row[1].value for row in wbx["Run Info"].iter_rows(min_row=2)}
ok("Excel download works and says who it was prepared for", xr.status_code == 200 and xr.data[:2] == b"PK"
   and info.get("Prepared for") == "Global Education Services")
ok("Excel keeps a formula as text", wbx["Data"]["C2"].data_type == "s")
ok("only csv and xlsx can be downloaded", cust.get(f"/s/{eid}/file.exe", headers=H2).status_code == 404)

mine = c.get("/api/shares?box=by-me").get_json()
row = next(x for x in mine if x["id"] == eid)
ok("the sender sees how many times it was opened and downloaded", row["views"] == 1 and row["downloads"] == 2
   and row["last_opened"] and row["first_opened"], str(row))
ok("another member cannot open the customer share inside the app", meena.get(f"/api/shares/{eid}").status_code == 404)
ok("the owner can still preview it inside the app", c.get(f"/api/shares/{eid}").status_code == 200)
ok("customer shares are not listed as Shared with me", all(x["id"] != eid for x in meena.get("/api/shares").get_json()))

# no passcode: counted for real visitors, not for the sender or chat-app link previews
open_ = c.post("/api/shares", json=cpay(name="No passcode", customer="Acme Edu")).get_json()
ok("a link without a passcode opens straight away", app.test_client().get(f"/s/{open_['id']}").status_code == 200)
app.test_client().get(f"/s/{open_['id']}", headers={"User-Agent": "WhatsApp/2.23.20 A"})
app.test_client().get(f"/s/{open_['id']}", headers={"User-Agent": "TelegramBot (like TwitterBot)"})
c.get(f"/s/{open_['id']}")
ok("chat-app link previews and the sender's own visits are not counted",
   next(x for x in c.get("/api/shares?box=by-me").get_json() if x["id"] == open_["id"])["views"] == 1)

prot = c.post("/api/shares", json=cpay(passcode="Owner123", name="Owner preview")).get_json()
ok("the sender can preview a protected customer page without the passcode",
   b"Owner preview" in c.get(f"/s/{prot['id']}").data and c.get(f"/s/{prot['id']}/file.csv").status_code == 200)
ok("the sender's preview is not counted as a customer visit",
   next(x for x in c.get("/api/shares?box=by-me").get_json() if x["id"] == prot["id"])["views"] == 0)
ok("someone else signed in still needs the passcode", b"protected" in meena.get(f"/s/{prot['id']}").data)

# who can reach what
gone = app.test_client().get("/s/doesnotexist123")
ok("an unknown link shows a friendly page, not an error", gone.status_code == 404 and b"no longer available" in gone.data)
internal = c.post("/api/shares", json=payload(link_access=True)).get_json()
ok("a colleague share is not reachable through the public address", app.test_client().get(f"/s/{internal['id']}").status_code == 404
   and app.test_client().get(f"/s/{internal['id']}/file.csv").status_code == 404)
admin_list = c.get("/api/admin/shares").get_json()
ok("an admin can list everything shared, without the data", any(x["id"] == eid for x in admin_list)
   and all("rows_data" not in x for x in admin_list))
ok("members cannot use the admin share list", meena.get("/api/admin/shares").status_code == 403)
ok("an admin can stop a customer link", meena.delete(f"/api/shares/{eid}").status_code == 404 and c.delete(f"/api/shares/{eid}").status_code == 200)
ok("a stopped link stops working at once", cust.get(f"/s/{eid}", headers=H2).status_code == 404
   and cust.get(f"/s/{eid}/file.csv", headers=H2).status_code == 404)
with app.app_context():
    row = _db.session.get(Share, open_["id"]); row.expires_at = _dt.utcnow() - _td(minutes=1); _db.session.commit()
ok("an expired link stops working", app.test_client().get(f"/s/{open_['id']}").status_code == 404)

# ---- database fallback: a dead database must not take the site down ---------------------------
os.environ["DATABASE_URL"] = "postgresql://user:pw@127.0.0.1:1/none"
t0 = time.time()
url, state = appmod._resolve_database()
ok("unreachable database falls back to SQLite", url.startswith("sqlite") and state["fallback"] and state["error"], f"{url} {state}")
ok("fallback is quick", time.time() - t0 < 15)
os.environ["DATABASE_URL"] = "postgres://u:p@host/db"
ok("Render-style postgres:// URLs use the psycopg driver", appmod._database_url().startswith("postgresql+psycopg://u:p@host/db"))

# ---- directory sources: saved per person, validated, presets, and never visible to anyone else -------------------
GOOD = {"name": "Hyderabad schools", "category": "Schools", "config": {
    "mode": "jsonld", "list_urls": ["https://www.example-directory.test/schools/hyderabad"], "pages": 2, "profile": True}}
r = c.post("/api/sources", json=GOOD)
ok("a person can save a source (structured-data mode, two pages, profile reading)", r.status_code == 200, r.get_data(as_text=True)[:200])
sid = r.get_json()["id"]
ok("the saved source comes back with its summary", r.get_json()["lists"] == 1 and r.get_json()["pages"] == 2)
ok("the owner sees their own sources", any(x["id"] == sid for x in c.get("/api/sources").get_json()["sources"]))
ok("another person does not see it", not any(x["id"] == sid for x in meena.get("/api/sources").get_json()["sources"]))
ok("another person cannot remove it", meena.delete(f"/api/sources/{sid}").status_code == 404)
ok("a source with no name is refused", c.post("/api/sources", json={**GOOD, "name": ""}).status_code == 400)
ok("a list address that is not a web address is refused",
   c.post("/api/sources", json={"name": "x", "config": {"mode": "jsonld", "list_urls": ["ftp://bad"]}}).status_code == 400)
ok("CSS mode needs its selectors",
   c.post("/api/sources", json={"name": "x", "config": {"mode": "css", "list_urls": ["https://a.test/"]}}).status_code == 400)
ok("a broken CSS selector is refused with a message, not saved",
   "not valid CSS" in c.post("/api/sources", json={"name": "x", "config": {"mode": "css", "list_urls": ["https://a.test/"],
       "item_selector": "div[[[", "name_selector": "h3"}}).get_json()["error"])
ok("pages above the limit are refused",
   c.post("/api/sources", json={"name": "x", "config": {"mode": "jsonld", "list_urls": ["https://a.test/"], "pages": 999}}).status_code == 400)
pres = c.get("/api/sources/presets").get_json()["presets"]
ok("the presets ship as data: colleges and schools are both there", {p["id"] for p in pres} >= {"colleges9-telangana-engineering", "edzy-schools-hyderabad"})
ok("a preset is copied into the person's own sources", c.post("/api/sources/presets/edzy-schools-hyderabad").status_code == 200)
ok("deleting a source removes it", c.delete(f"/api/sources/{sid}").status_code == 200
   and not any(x["id"] == sid for x in c.get("/api/sources").get_json()["sources"]))
import scraper.fetch as _fetch_mod
class _FakeSourceFetcher:
    def __init__(self, spec=None): pass
    def get_html(self, url):
        if url.endswith("/hyderabad"):
            return ('<script type="application/ld+json">{"@type":"ItemList","itemListElement":[{"@type":"School","name":"Test School","url":"https://www.example-directory.test/s/1"}]}</script>', "ok")
        return ('<html><body>Phone No. 040-12345678 Email info@testschool.in</body></html>', "ok")
appmod.Fetcher = _FakeSourceFetcher
t = c.post("/api/sources/test", json={"config": GOOD["config"]}).get_json()
ok("the test reads the first page and the first entry's own page", t["ok"] and t["count"] == 1 and t["sample"][0]["name"] == "Test School"
   and t["profile"] and "040-12345678" in t["profile"]["phones"][0], t)
appmod.Fetcher = _fetch_mod.Fetcher


print(f"\n{len(failures)} failure(s)" if failures else "\nAll tests passed")
sys.exit(1 if failures else 0)
