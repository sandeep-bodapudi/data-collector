"""Write results to a formatted Excel workbook."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
# Settings that must never appear in an exported file.
SECRET_KEYS = {"ai_api_key", "ai_key", "brave_key", "li_at_cookie", "fb_cookie", "ai", "allow_restricted"}
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")


def _cell_value(v):
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    if v is None:
        return ""
    if isinstance(v, str) and len(v) > 32000:  # Excel cell limit is 32,767 chars
        return v[:32000] + "…"
    return v


def write_workbook(path: str, columns: list[str], rows: list[dict], run_info: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(columns)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        ws.append([_cell_value(row.get(col)) for col in columns])
    # Scraped text that starts with "=" must stay text, never become a live Excel formula.
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        longest = max([len(str(col))] + [len(str(ws.cell(r, idx).value or "")) for r in range(2, min(ws.max_row, 200) + 1)])
        ws.column_dimensions[letter].width = max(12, min(60, longest + 2))
        is_link = "url" in col.lower() or "link" in col.lower() or col in ("Website", "Source URL")
        if is_link:
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(r, idx)
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = LINK_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Run Info")
    info.append(["Setting", "Value"])
    for c in info[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    info.append(["Created", datetime.now().strftime("%Y-%m-%d %H:%M")])
    info.append(["Total rows", len(rows)])
    for k, v in run_info.items():
        if k not in SECRET_KEYS:
            info.append([k.replace("_", " ").capitalize(), _cell_value(v)])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 90

    wb.save(path)
