"""Reads uploaded files (xlsx, csv, tsv, json) into rows. Stateless: nothing is stored on the server."""
import csv
import io
import json
import os
from datetime import date, datetime, time

from openpyxl import load_workbook


MAX_IMPORT_ROWS = 200_000


def _cell(v):
    """JSON-safe cell: dates become ISO text, everything else numeric/text stays as is."""
    if v is None:
        return ""
    if isinstance(v, (datetime, date, time)):
        return v.isoformat(sep=" ") if isinstance(v, datetime) else v.isoformat()
    if isinstance(v, (bool, int, float, str)):
        return v
    return str(v)


def _xlsx_rows(data: bytes) -> tuple[list[str], list[dict]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = next(it, None) or []
    columns, seen = [], set()
    for i, h in enumerate(header):  # blank or repeated headers get a unique name
        name = str(h).strip() if h not in (None, "") else f"Column {i + 1}"
        while name in seen:
            name += " (2)"
        seen.add(name)
        columns.append(name)
    rows = []
    for values in it:
        if values is None or all(v in (None, "") for v in values):
            continue
        rows.append({c: _cell(v) for c, v in zip(columns, values)})
        if len(rows) >= MAX_IMPORT_ROWS:
            break
    wb.close()
    return columns, rows


def parse_upload(filename: str, data: bytes) -> tuple[list[str], list[dict]]:
    ext = os.path.splitext(filename.lower())[1]
    if ext in (".xlsx", ".xlsm"):
        try:
            columns, rows = _xlsx_rows(data)
        except Exception:
            raise ValueError("This Excel file could not be read. Save it as .xlsx and try again.")
    elif ext in (".csv", ".tsv", ".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = "excel-tab" if ext == ".tsv" or text.split("\n", 1)[0].count("\t") > text.split("\n", 1)[0].count(",") else "excel"
        reader = csv.reader(io.StringIO(text), dialect=dialect)
        header = next(reader, [])
        columns = [h.strip() or f"Column {i + 1}" for i, h in enumerate(header)]
        rows = [dict(zip(columns, r)) for r in reader if any(c.strip() for c in r)][:MAX_IMPORT_ROWS]
    elif ext == ".json":
        try:
            items = json.loads(data.decode("utf-8-sig"))
        except ValueError:
            raise ValueError("This JSON file is not valid.")
        if isinstance(items, dict):
            items = next((v for v in items.values() if isinstance(v, list)), [items])
        rows = [i for i in items if isinstance(i, dict)][:MAX_IMPORT_ROWS]
        columns = list(dict.fromkeys(k for r in rows for k in r))
        rows = [{c: (json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v) for c, v in r.items()} for r in rows]
    else:
        raise ValueError("Use an .xlsx, .csv, .tsv or .json file.")
    if not columns or not rows:
        raise ValueError("The file has no rows to import.")
    return columns, rows
